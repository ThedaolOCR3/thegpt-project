import asyncio
import unittest
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pymupdf
from docx import Document as WordDocument
from fastapi import UploadFile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from starlette.datastructures import Headers

from app.services.hybrid_ocr.chunk_service import create_chunks
from app.services.hybrid_ocr.document_processing_service import (
    DocumentProcessingService,
)
from app.services.hybrid_ocr.errors import (
    DocumentTooLargeError,
    DocumentValidationError,
    OfficeExtractionError,
)
from app.services.hybrid_ocr.models import (
    OcrEngineResult,
    OcrProcessingConfig,
)


class FakeOcrEngine:
    """실제 모델 다운로드 없이 PaddleOCR 호출 횟수와 반환 흐름을 검증합니다."""

    def __init__(self, text: str = "가짜 OCR 추출 텍스트") -> None:
        self.text = text
        self.call_count = 0
        self.last_image_size: tuple[int, int] | None = None

    def extract_text(self, image: Image.Image) -> OcrEngineResult:
        self.call_count += 1
        self.last_image_size = image.size
        return OcrEngineResult(
            text=self.text,
            confidence=0.95 if self.text else 0.0,
            line_count=1 if self.text else 0,
            processing_time_seconds=0.01,
        )


class HybridOcrServiceTest(unittest.TestCase):
    def setUp(self) -> None:
        self.ocr_engine = FakeOcrEngine()
        self.service = DocumentProcessingService(
            config=OcrProcessingConfig(
                max_file_bytes=5 * 1024 * 1024,
                max_pdf_pages=10,
                native_text_min_chars=20,
                significant_image_area_ratio=0.03,
                pdf_render_dpi=100,
                max_image_side=1200,
                max_image_pixels=10_000_000,
                paddle_device="cpu",
                paddle_language="korean",
            ),
            ocr_service_factory=lambda: self.ocr_engine,
        )

    def test_digital_pdf_uses_native_text_without_ocr(self) -> None:
        pdf = _make_digital_pdf()

        result = self._process(pdf, "digital.pdf", "application/pdf")

        self.assertEqual(self.ocr_engine.call_count, 0)
        self.assertEqual(result.page_count, 1)
        self.assertIn("Native digital PDF text", result.extracted_text)
        self.assertIn("디지털 PDF", result.notes[0])

    def test_scanned_pdf_renders_page_and_calls_ocr_once(self) -> None:
        pdf = _make_scanned_pdf()

        result = self._process(pdf, "scanned.pdf", "application/pdf")

        self.assertEqual(self.ocr_engine.call_count, 1)
        self.assertIn("가짜 OCR 추출 텍스트", result.extracted_text)
        self.assertIn("스캔 PDF", result.notes[0])

    def test_hybrid_pdf_inserts_image_ocr_after_native_text(self) -> None:
        pdf = _make_hybrid_pdf()

        result = self._process(pdf, "hybrid.pdf", "application/pdf")

        self.assertEqual(self.ocr_engine.call_count, 1)
        self.assertIn("Native text before the embedded image", result.extracted_text)
        self.assertIn("[이미지 OCR]", result.extracted_text)
        self.assertIn("Hybrid PDF", result.notes[0])
        self.assertLess(
            result.extracted_text.index("Native text before"),
            result.extracted_text.index("[이미지 OCR]"),
        )
        self.assertLess(
            result.extracted_text.index("[이미지 OCR]"),
            result.extracted_text.index("Native text after"),
        )

    def test_png_calls_ocr_once(self) -> None:
        image = _make_png()

        result = self._process(image, "sample.png", "image/png")

        self.assertEqual(self.ocr_engine.call_count, 1)
        self.assertEqual(result.page_count, 1)
        self.assertEqual(result.confidence, 95.0)

    def test_jpg_is_supported(self) -> None:
        result = self._process(_make_jpg(), "sample.jpg", "image/jpeg")

        self.assertEqual(self.ocr_engine.call_count, 1)
        self.assertIn("가짜 OCR 추출 텍스트", result.extracted_text)

    def test_docx_extracts_heading_paragraph_and_table_without_ocr(self) -> None:
        result = self._process(
            _make_office_package("docx"),
            "sample.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        self.assertEqual(self.ocr_engine.call_count, 0)
        self.assertIn("DOCX direct extraction heading", result.extracted_text)
        self.assertIn("DOCX paragraph for RAG extraction", result.extracted_text)
        self.assertIn("| Field | Value |", result.extracted_text)
        self.assertIsNone(result.page_count)
        self.assertIn("원본 형식: DOCX", result.notes)
        self.assertTrue(any("실제 페이지 수" in note for note in result.notes))

    def test_pptx_extracts_slide_text_and_notes_without_ocr(self) -> None:
        result = self._process(
            _make_office_package("pptx"),
            "slides.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

        self.assertEqual(self.ocr_engine.call_count, 0)
        self.assertEqual(result.page_count, 1)
        self.assertIn("PPTX direct extraction title", result.extracted_text)
        self.assertIn("PPTX speaker notes", result.extracted_text)
        self.assertIn("원본 형식: PPTX", result.notes)

    def test_docx_embedded_image_is_sent_to_ocr(self) -> None:
        result = self._process(
            _make_office_package("docx", include_image=True),
            "image.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        self.assertEqual(self.ocr_engine.call_count, 1)
        self.assertIn("### 이미지 OCR 1", result.extracted_text)
        self.assertIn("가짜 OCR 추출 텍스트", result.extracted_text)

    def test_pptx_embedded_image_is_sent_to_ocr(self) -> None:
        result = self._process(
            _make_office_package("pptx", include_image=True),
            "image.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

        self.assertEqual(self.ocr_engine.call_count, 1)
        self.assertIn("가짜 OCR 추출 텍스트", result.extracted_text)

    def test_large_image_is_resized_before_ocr(self) -> None:
        image = _make_png(width=2000, height=1000)

        self._process(image, "large.png", "image/png")

        self.assertEqual(self.ocr_engine.last_image_size, (1200, 600))

    def test_exif_orientation_is_applied_before_ocr(self) -> None:
        image = _make_rotated_jpg()

        self._process(image, "rotated.jpg", "image/jpeg")

        self.assertEqual(self.ocr_engine.last_image_size, (200, 400))

    def test_multi_page_digital_pdf_keeps_page_order(self) -> None:
        pdf = _make_multi_page_digital_pdf()

        result = self._process(pdf, "multi.pdf", "application/pdf")

        self.assertEqual(self.ocr_engine.call_count, 0)
        self.assertEqual(result.page_count, 2)
        self.assertLess(
            result.extracted_text.index("First page native text"),
            result.extracted_text.index("Second page native text"),
        )

    def test_empty_ocr_result_requires_review(self) -> None:
        self.ocr_engine.text = ""

        result = self._process(_make_scanned_pdf(), "empty.pdf", "application/pdf")

        self.assertEqual(result.readiness, "review")
        self.assertEqual(result.confidence, 0.0)

    def test_rejects_corrupted_pdf(self) -> None:
        with self.assertRaises(DocumentValidationError):
            self._process(b"%PDF-corrupted", "broken.pdf", "application/pdf")

    def test_rejects_encrypted_pdf(self) -> None:
        with self.assertRaises(DocumentValidationError):
            self._process(_make_encrypted_pdf(), "locked.pdf", "application/pdf")

    def test_rejects_corrupted_image(self) -> None:
        with self.assertRaises(DocumentValidationError):
            self._process(b"not-an-image", "broken.png", "image/png")

    def test_rejects_wrong_mime_type(self) -> None:
        with self.assertRaises(DocumentValidationError):
            self._process(_make_png(), "sample.png", "text/plain")

    def test_rejects_file_over_size_limit(self) -> None:
        small_limit_service = DocumentProcessingService(
            config=OcrProcessingConfig(
                **{
                    **self.service.config.__dict__,
                    "max_file_bytes": 10,
                }
            ),
            ocr_service_factory=lambda: self.ocr_engine,
        )

        with self.assertRaises(DocumentTooLargeError):
            asyncio.run(
                small_limit_service.process_document(
                    file=_upload(_make_png(), "large.png", "image/png"),
                    chunk_size=200,
                    overlap=30,
                )
            )

    def test_rejects_extension_and_binary_mismatch(self) -> None:
        with self.assertRaises(DocumentValidationError):
            self._process(_make_png(), "not-really.pdf", "application/pdf")

    def test_rejects_corrupted_office_archive(self) -> None:
        with self.assertRaises(DocumentValidationError):
            self._process(
                b"PK-not-a-valid-archive",
                "broken.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

    def test_rejects_docx_renamed_as_pptx(self) -> None:
        with self.assertRaises(DocumentValidationError):
            self._process(
                _make_office_package("docx"),
                "renamed.pptx",
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )

    def test_rejects_wrong_office_mime_type(self) -> None:
        with self.assertRaises(DocumentValidationError):
            self._process(
                _make_office_package("docx"),
                "sample.docx",
                "text/plain",
            )

    def test_valid_zip_with_unreadable_docx_structure_raises_domain_error(self) -> None:
        with self.assertRaises(OfficeExtractionError):
            self._process(
                _make_minimal_unreadable_office_package("docx"),
                "unreadable.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

    def test_rejects_overlap_equal_to_chunk_size(self) -> None:
        upload = _upload(_make_png(), "sample.png", "image/png")

        with self.assertRaises(DocumentValidationError):
            asyncio.run(
                self.service.process_document(
                    file=upload,
                    chunk_size=100,
                    overlap=100,
                )
            )

    def test_chunking_moves_forward_and_keeps_overlap(self) -> None:
        text = "문서 처리 흐름을 확인하기 위한 문장입니다. " * 20

        chunks = create_chunks(text, chunk_size=120, overlap=20)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk for chunk in chunks))
        self.assertTrue(all(len(chunk) <= 120 for chunk in chunks))

    def _process(
        self,
        content: bytes,
        file_name: str,
        content_type: str,
    ):
        return asyncio.run(
            self.service.process_document(
                file=_upload(content, file_name, content_type),
                chunk_size=200,
                overlap=30,
            )
        )


def _upload(content: bytes, file_name: str, content_type: str) -> UploadFile:
    return UploadFile(
        file=BytesIO(content),
        filename=file_name,
        headers=Headers({"content-type": content_type}),
    )


def _make_office_package(file_type: str, include_image: bool = False) -> bytes:
    output = BytesIO()
    if file_type == "docx":
        document = WordDocument()
        document.add_heading("DOCX direct extraction heading", level=1)
        document.add_paragraph("DOCX paragraph for RAG extraction")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Field"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Diagnosis"
        table.cell(1, 1).text = "Sample"
        if include_image:
            document.add_picture(BytesIO(_make_png()), width=Inches(2))
        document.save(output)
    elif file_type == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        text_box.text = "PPTX direct extraction title"
        slide.notes_slide.notes_text_frame.text = "PPTX speaker notes"
        if include_image:
            slide.shapes.add_picture(
                BytesIO(_make_png()),
                Inches(1),
                Inches(2),
                width=Inches(4),
            )
        presentation.save(output)
    else:
        raise ValueError(f"지원하지 않는 테스트 Office 형식: {file_type}")
    return output.getvalue()


def _make_minimal_unreadable_office_package(file_type: str) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?><Types '
            'xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8"?><Relationships '
            'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>',
        )
        if file_type == "docx":
            archive.writestr(
                "word/document.xml",
                '<?xml version="1.0" encoding="UTF-8"?><w:document '
                'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                "<w:body/></w:document>",
            )
        elif file_type == "pptx":
            archive.writestr(
                "ppt/presentation.xml",
                '<?xml version="1.0" encoding="UTF-8"?><p:presentation '
                'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>',
            )
        else:
            raise ValueError(f"지원하지 않는 테스트 Office 형식: {file_type}")
    return output.getvalue()


def _make_png(width: int = 500, height: int = 300) -> bytes:
    output = BytesIO()
    Image.new("RGB", (width, height), "white").save(output, format="PNG")
    return output.getvalue()


def _make_jpg(width: int = 500, height: int = 300) -> bytes:
    output = BytesIO()
    Image.new("RGB", (width, height), "white").save(output, format="JPEG")
    return output.getvalue()


def _make_rotated_jpg() -> bytes:
    output = BytesIO()
    image = Image.new("RGB", (400, 200), "white")
    exif = Image.Exif()
    exif[274] = 6
    image.save(output, format="JPEG", exif=exif)
    return output.getvalue()


def _make_digital_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Native digital PDF text with enough characters for direct extraction.",
    )
    content = document.tobytes()
    document.close()
    return content


def _make_scanned_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=500, height=300)
    page.insert_image(page.rect, stream=_make_png())
    content = document.tobytes()
    document.close()
    return content


def _make_multi_page_digital_pdf() -> bytes:
    document = pymupdf.open()
    first_page = document.new_page()
    first_page.insert_text(
        (72, 72),
        "First page native text with enough characters for direct extraction.",
    )
    second_page = document.new_page()
    second_page.insert_text(
        (72, 72),
        "Second page native text with enough characters for direct extraction.",
    )
    content = document.tobytes()
    document.close()
    return content


def _make_encrypted_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Password protected PDF")
    content = document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner-password",
        user_pw="user-password",
    )
    document.close()
    return content


def _make_hybrid_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=600, height=800)
    page.insert_text(
        (50, 60),
        "Native text before the embedded image with enough characters to use hybrid mode.",
    )
    page.insert_image(pymupdf.Rect(50, 100, 550, 400), stream=_make_png())
    page.insert_text(
        (50, 450),
        "Native text after the embedded image keeps the original document order.",
    )
    content = document.tobytes()
    document.close()
    return content


if __name__ == "__main__":
    unittest.main()
